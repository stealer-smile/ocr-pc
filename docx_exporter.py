"""
docx_exporter.py — Convert Markdown/HTML OCR output to formatted DOCX.

Strategy: MD → HTML (via markdown lib) → parse HTML → DOCX (via python-docx).
This handles both markdown tables AND raw HTML tables from OCR output.
"""
import re
from html.parser import HTMLParser
from pathlib import Path

import markdown
from docx import Document
from docx.shared import Pt, Inches, RGBColor, Emu, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement


class _DocxBuilder(HTMLParser):
  """HTML parser that builds a python-docx Document."""

  def __init__(self, doc: Document):
    super().__init__()
    self.doc = doc
    self._p = None          # current paragraph
    self._bold = False
    self._italic = False
    self._code = False
    self._heading = 0       # 0 = not heading, 1-6 = h1-h6
    self._in_pre = False
    self._pre_lines: list[str] = []

    # Table state
    self._in_table = False
    self._table = None
    self._row_idx = -1
    self._col_idx = -1
    self._cell_text = ""
    self._is_th = False
    self._table_rows: list[list[str]] = []
    self._current_row: list[str] = []

    # List state
    self._in_li = False
    self._li_text = ""
    self._list_type: list[str] = []  # stack: "ul" or "ol"

    self._pending_text = ""

  def _ensure_paragraph(self):
    if self._p is None:
      self._p = self.doc.add_paragraph()
    return self._p

  def _flush_paragraph(self):
    self._p = None

  def _add_run(self, text: str):
    if not text:
      return
    p = self._ensure_paragraph()
    run = p.add_run(text)
    if self._bold:
      run.bold = True
    if self._italic:
      run.italic = True
    if self._code and not self._in_pre:
      run.font.name = "Consolas"
      run.font.size = Pt(11)

  def handle_starttag(self, tag: str, attrs: list):
    tag = tag.lower()

    if tag in ("h1", "h2", "h3", "h4", "h5", "h6"):
      self._flush_paragraph()
      level = int(tag[1])
      self._p = self.doc.add_heading(level=level)
      self._heading = level

    elif tag == "strong" or tag == "b":
      self._bold = True
    elif tag == "em" or tag == "i":
      self._italic = True
    elif tag == "code":
      self._code = True
    elif tag == "pre":
      self._in_pre = True
      self._pre_lines = []

    elif tag == "br":
      if self._in_table:
        self._cell_text += "\n"
      elif self._in_li:
        self._li_text += "\n"
      else:
        self._add_run("\n")

    elif tag == "hr":
      self._flush_paragraph()
      p = self.doc.add_paragraph()
      pPr = p.paragraph_format.element.get_or_add_pPr()
      pBdr = OxmlElement("w:pBdr")
      bottom = OxmlElement("w:bottom")
      bottom.set(qn("w:val"), "single")
      bottom.set(qn("w:sz"), "6")
      bottom.set(qn("w:space"), "1")
      bottom.set(qn("w:color"), "999999")
      pBdr.append(bottom)
      pPr.append(pBdr)
      self._flush_paragraph()

    elif tag == "p":
      self._flush_paragraph()

    elif tag == "table":
      self._flush_paragraph()
      self._in_table = True
      self._table_rows = []

    elif tag in ("thead", "tbody", "tfoot"):
      pass  # ignore, we track rows directly

    elif tag == "tr":
      self._current_row = []

    elif tag in ("td", "th"):
      self._is_th = (tag == "th")
      self._cell_text = ""

    elif tag in ("ul", "ol"):
      self._list_type.append(tag)

    elif tag == "li":
      self._in_li = True
      self._li_text = ""

  def handle_endtag(self, tag: str):
    tag = tag.lower()

    if tag in ("h1", "h2", "h3", "h4", "h5", "h6"):
      self._heading = 0
      self._flush_paragraph()

    elif tag == "strong" or tag == "b":
      self._bold = False
    elif tag == "em" or tag == "i":
      self._italic = False
    elif tag == "code":
      if self._in_pre:
        pass  # will flush in </pre>
      self._code = False

    elif tag == "pre":
      # Flush code block
      self._flush_paragraph()
      code = "\n".join(self._pre_lines)
      p = self.doc.add_paragraph()
      run = p.add_run(code)
      run.font.name = "Consolas"
      run.font.size = Pt(11)
      # Gray background
      shading = OxmlElement("w:shd")
      shading.set(qn("w:val"), "clear")
      shading.set(qn("w:color"), "auto")
      shading.set(qn("w:fill"), "F0F0F0")
      p.paragraph_format.element.get_or_add_pPr().append(shading)
      self._in_pre = False
      self._pre_lines = []
      self._flush_paragraph()

    elif tag == "p":
      self._flush_paragraph()

    elif tag in ("td", "th"):
      self._current_row.append(self._cell_text.strip())

    elif tag == "tr":
      self._table_rows.append(self._current_row)
      self._current_row = []

    elif tag == "table":
      # Build DOCX table from collected rows
      if self._table_rows:
        n_rows = len(self._table_rows)
        n_cols = max(len(r) for r in self._table_rows) if self._table_rows else 1
        tbl = self.doc.add_table(rows=n_rows, cols=n_cols)
        tbl.style = "Table Grid"
        for ri, row in enumerate(self._table_rows):
          for ci, cell_text in enumerate(row):
            if ci < n_cols:
              cell = tbl.rows[ri].cells[ci]
              cell.text = cell_text
              # Bold first row (header)
              if ri == 0:
                for paragraph in cell.paragraphs:
                  for run in paragraph.runs:
                    run.bold = True
      self._in_table = False
      self._table_rows = []
      self._flush_paragraph()

    elif tag in ("thead", "tbody", "tfoot"):
      pass

    elif tag == "li":
      self._in_li = False
      self._flush_paragraph()
      style = "List Number" if (self._list_type and self._list_type[-1] == "ol") else "List Bullet"
      p = self.doc.add_paragraph(style=style)
      p.add_run(self._li_text.strip())
      self._flush_paragraph()

    elif tag in ("ul", "ol"):
      if self._list_type:
        self._list_type.pop()

  def handle_data(self, data: str):
    if self._in_pre:
      self._pre_lines.append(data.rstrip("\n"))
      return

    if self._in_table:
      self._cell_text += data
      return

    if self._in_li:
      self._li_text += data
      return

    # Normal text
    text = data
    if not self._heading and not self._code:
      # Collapse whitespace for normal paragraphs
      text = re.sub(r'\n+', ' ', text)

    self._add_run(text)


def md_to_docx(md_text: str, output_path: str) -> None:
  """Convert markdown (possibly with embedded HTML) to a formatted DOCX file."""
  # Step 1: markdown → HTML
  html = markdown.markdown(
    md_text,
    extensions=["tables", "fenced_code", "nl2br"],
  )

  # Step 2: parse HTML → DOCX
  doc = Document()

  # ── Page setup: A4, margins chuẩn (top 2, bottom 2, left 3, right 2 cm) ──
  for section in doc.sections:
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)
    section.left_margin = Cm(3)
    section.right_margin = Cm(2)

  # ── Default style: Times New Roman 13pt, justified, line spacing 1.5 ──
  style = doc.styles["Normal"]
  style.font.name = "Times New Roman"
  style.font.size = Pt(13)
  style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
  style.paragraph_format.line_spacing = 1.5
  style.paragraph_format.space_before = Pt(0)
  style.paragraph_format.space_after = Pt(6)
  style.paragraph_format.first_line_indent = Cm(1.27)

  # ── Heading styles: Times New Roman, bold ──
  for i in range(1, 7):
    hs = doc.styles[f"Heading {i}"]
    hs.font.name = "Times New Roman"
    hs.font.color.rgb = RGBColor(0, 0, 0)
    hs.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    hs.paragraph_format.space_before = Pt(12)
    hs.paragraph_format.space_after = Pt(6)
    hs.paragraph_format.first_line_indent = Cm(0)

  # Heading sizes
  doc.styles["Heading 1"].font.size = Pt(16)
  doc.styles["Heading 2"].font.size = Pt(14)
  doc.styles["Heading 3"].font.size = Pt(13)

  # ── List styles ──
  for ls_name in ("List Bullet", "List Number"):
    if ls_name in doc.styles:
      ls = doc.styles[ls_name]
      ls.font.name = "Times New Roman"
      ls.font.size = Pt(13)
      ls.paragraph_format.first_line_indent = Cm(0)

  builder = _DocxBuilder(doc)
  builder.feed(html)

  doc.save(output_path)
